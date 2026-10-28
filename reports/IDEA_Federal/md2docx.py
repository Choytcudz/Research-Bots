import re
from docx import Document
from docx.shared import Pt
src=open('IDEA_Federal_Process_Report_PRISMA_2026-10-28.md').read().splitlines()
d=Document(); st=d.styles['Normal']; st.font.name='Times New Roman'; st.font.size=Pt(11)
def runs(p,text):
    for part in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)',text):
        if part.startswith('**') and part.endswith('**') and len(part)>4: p.add_run(part[2:-2]).bold=True
        elif part.startswith('*') and part.endswith('*') and len(part)>2: p.add_run(part[1:-1]).italic=True
        else: p.add_run(part)
code=False; table=None
for line in src:
    if line.startswith('```'):
        code=not code; continue
    if code:
        p=d.add_paragraph(); r=p.add_run(line); r.font.name='Courier New'; r.font.size=Pt(8)
        p.paragraph_format.space_after=Pt(0); continue
    if line.startswith('|'):
        cells=[c.strip() for c in line.strip('|').split('|')]
        if all(re.fullmatch(r'-+',c) for c in cells): continue
        if table is None:
            table=d.add_table(rows=0,cols=len(cells)); table.style='Table Grid'
        row=table.add_row().cells
        for i,c in enumerate(cells[:len(row)]): row[i].text=c.replace('**','')
        continue
    table=None
    m=re.match(r'(#+) (.*)',line)
    if m: d.add_heading(m.group(2),level=min(len(m.group(1)),3)); continue
    if line.startswith('- '): runs(d.add_paragraph(style='List Bullet'),line[2:]); continue
    if re.match(r'^\d+\. ',line): runs(d.add_paragraph(style='List Number'),re.sub(r'^\d+\. ','',line)); continue
    if line.strip(): runs(d.add_paragraph(),line)
cp=d.core_properties; cp.author='Joseph Castellanos Hoyt'; cp.last_modified_by='Joseph Castellanos Hoyt'; cp.title='IDEA Federal (federal IDEA baseline reviewer, Federal) Process Report'
import datetime; cp.created=cp.modified=datetime.datetime(2026,10,28,22,0,0)
d.save('IDEA_Federal_Process_Report_PRISMA_2026-10-28.docx')
